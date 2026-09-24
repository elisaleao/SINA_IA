from unittest.mock import AsyncMock, MagicMock

import docx
import pymupdf as fitz
import pytest
from PIL import Image

from app.services.document_extractor import DocumentExtractor
from app.services.ingestion_service import (
    INGESTION_OCR_PROMPT,
    IngestionService,
)


@pytest.mark.asyncio
async def test_process_pdf_with_text_and_equations(tmp_path):
    pdf_path = tmp_path / 'documento_teste.pdf'

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 72),
        'Introdução à Física Quântica: A energia é calculada por $$E = m c^2$$. '
        'Além disso, $x + y = 10$ representa uma reta no plano.',
    )
    doc.save(str(pdf_path))
    doc.close()

    service = _service(_fake_gemini())
    markdown, accessible, equations = await service.process_file(
        str(pdf_path), 'documento_teste.pdf'
    )

    assert '## Página 1' in markdown
    assert 'Física Quântica' in markdown
    assert len(equations) >= 2
    assert any('E = m c^2' in eq for eq in equations)
    assert any('x + y = 10' in eq for eq in equations)
    assert 'm c ao quadrado' in accessible


@pytest.mark.asyncio
async def test_process_docx_headings_and_paragraphs(tmp_path):
    docx_path = tmp_path / 'aula_algoritmos.docx'

    doc = docx.Document()
    doc.add_heading('Estruturas de Dados', level=1)
    doc.add_heading('Pilhas e Filas', level=2)
    doc.add_paragraph(
        'Uma pilha segue a política LIFO com complexidade $O(1)$.'
    )
    doc.save(str(docx_path))

    service = _service(_fake_gemini())
    markdown, accessible, equations = await service.process_file(
        str(docx_path), 'aula_algoritmos.docx'
    )

    assert '# Estruturas de Dados' in markdown
    assert '## Pilhas e Filas' in markdown
    assert 'complexidade' in markdown
    assert 'O(1)' in equations
    assert isinstance(accessible, str)


@pytest.mark.asyncio
async def test_process_plain_text(tmp_path):
    txt_path = tmp_path / 'notas.txt'
    txt_path.write_text(
        'Conteúdo em texto puro com fórmula: $a^2 + b^2 = c^2$.',
        encoding='utf-8',
    )

    service = _service(_fake_gemini())
    markdown, accessible, equations = await service.process_file(
        str(txt_path), 'notas.txt'
    )

    assert 'Conteúdo em texto puro' in markdown
    assert 'a^2 + b^2 = c^2' in equations
    assert isinstance(accessible, str)


def _png_file(path, color=(255, 255, 255)) -> str:
    Image.new('RGB', (60, 60), color).save(path, format='PNG')
    return str(path)


def _service(gemini) -> IngestionService:
    extractor = DocumentExtractor(
        gemini, max_visual_candidates=0, ocr_prompt=INGESTION_OCR_PROMPT
    )
    return IngestionService(gemini=gemini, extractor=extractor)


def _fake_gemini(ocr_text: str | None = None) -> MagicMock:
    gemini = MagicMock()
    gemini.is_configured = ocr_text is not None
    gemini.ocr_image = AsyncMock(return_value=ocr_text or '')
    return gemini


@pytest.mark.asyncio
async def test_process_image_without_client(tmp_path):
    img_path = _png_file(tmp_path / 'exemplo.png')
    gemini = _fake_gemini(ocr_text=None)
    service = _service(gemini)

    markdown, accessible, equations = await service.process_file(
        img_path, 'exemplo.png'
    )

    assert 'indisponível' in markdown.lower()
    assert equations == []
    assert isinstance(accessible, str)
    gemini.ocr_image.assert_not_called()


@pytest.mark.asyncio
async def test_process_image_with_mocked_gemini(tmp_path):
    img_path = _png_file(tmp_path / 'esquema.png', color=(128, 128, 128))
    gemini = _fake_gemini(
        'Transcrição OCR: Fórmula quadrática '
        '$x = \\frac{-b \\pm \\sqrt{\\Delta}}{2a}$.'
    )
    service = _service(gemini)

    markdown, accessible, equations = await service.process_file(
        img_path, 'esquema.png'
    )

    assert 'Transcrição OCR' in markdown
    assert len(equations) > 0
    assert isinstance(accessible, str)
    # O OCR recebe o prompt que pede fórmulas em LaTeX delimitado
    assert gemini.ocr_image.await_args.kwargs['prompt'] == INGESTION_OCR_PROMPT


@pytest.mark.asyncio
async def test_process_scanned_pdf_page_triggers_gemini_ocr(tmp_path):
    pdf_path = tmp_path / 'digitalizado.pdf'

    # Página em branco: sem texto nativo, dispara o OCR
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    gemini = _fake_gemini('Texto recuperado via OCR da imagem da página 1.')
    service = _service(gemini)

    markdown, accessible, equations = await service.process_file(
        str(pdf_path), 'digitalizado.pdf'
    )

    assert '## Página 1' in markdown
    assert 'Texto recuperado via OCR' in markdown
    gemini.ocr_image.assert_awaited_once()


@pytest.mark.asyncio
async def test_scanned_pdf_without_client_keeps_native_text(tmp_path):
    pdf_path = tmp_path / 'digitalizado.pdf'
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    gemini = _fake_gemini(ocr_text=None)
    service = _service(gemini)

    markdown, _, _ = await service.process_file(
        str(pdf_path), 'digitalizado.pdf'
    )

    assert markdown == '## Página 1'
    gemini.ocr_image.assert_not_called()


@pytest.mark.asyncio
async def test_unsupported_extension_is_rejected(tmp_path):
    csv_path = tmp_path / 'notas.csv'
    csv_path.write_text('a,b\n1,2', encoding='utf-8')

    with pytest.raises(ValueError, match='Formato não suportado'):
        await _service(_fake_gemini()).process_file(str(csv_path), 'notas.csv')
