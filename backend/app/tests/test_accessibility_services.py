from __future__ import annotations

import os
import time
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import fitz
import pytest
from docx import Document

from app.core.config import settings
from app.schemas.accessibility import (
    AuditItem,
    AuditReport,
    ChartVisionResult,
)
from app.services import gemini_service as gemini_module
from app.services import tts_service as tts_module
from app.services.accessibility_pipeline import AccessibilityPipeline
from app.services.document_extractor import (
    DocumentExtractor,
    ExtractionResult,
    VisualCandidate,
)
from app.services.file_store import GeneratedFileStore
from app.services.gemini_service import GeminiService
from app.services.tts_service import TTSService


class FakeOCRAI:
    async def ocr_image(
        self, image_bytes: bytes, mime_type: str, prompt: str = ''
    ) -> str:
        assert image_bytes
        assert mime_type.startswith('image/')
        return 'OCR extracted text'


class FakeTTS:
    async def synthesize(self, text: str, output_path: Path) -> None:
        assert text
        output_path.write_bytes(b'fake-mp3')


class FakeExtractor:
    def __init__(self, result: ExtractionResult) -> None:
        self.result = result

    async def extract(self, filename: str, data: bytes) -> ExtractionResult:
        assert filename
        assert data
        return self.result


class FakePipelineAI:
    def __init__(
        self, *, audit: AuditReport, corrected: str | None = None
    ) -> None:
        self.audit_result = audit
        self.corrected = corrected
        self.last_system_instruction = ''
        self.last_user_text = ''

    async def generate_text(
        self,
        *,
        system_instruction: str,
        user_text: str,
        temperature: float = 0.15,
    ) -> str:
        assert temperature == 0.15
        self.last_system_instruction = system_instruction
        self.last_user_text = user_text
        return 'Accessible generated text'

    async def analyze_chart(
        self,
        *,
        image_bytes: bytes,
        mime_type: str,
        context_text: str,
    ) -> ChartVisionResult:
        assert image_bytes
        assert mime_type == 'image/png'
        assert context_text
        return ChartVisionResult(
            is_chart=True,
            title='Sales chart',
            description='The chart shows a rising trend.',
            confidence=0.95,
        )

    async def audit(self, original: str, accessible: str) -> AuditReport:
        assert original
        assert accessible
        return self.audit_result

    async def correct(self, accessible: str, problems: list[str]) -> str:
        assert accessible
        assert problems
        return self.corrected or accessible


def make_pipeline(
    *,
    tmp_path: Path,
    extraction: ExtractionResult,
    ai: FakePipelineAI,
    monkeypatch: pytest.MonkeyPatch,
) -> AccessibilityPipeline:
    monkeypatch.setattr(
        settings, 'ACCESSIBILITY_OUTPUT_DIR', str(tmp_path / 'generated')
    )
    pipeline = AccessibilityPipeline.__new__(AccessibilityPipeline)
    pipeline.ai = ai
    pipeline.extractor = FakeExtractor(extraction)
    pipeline.tts = FakeTTS()
    pipeline.store = GeneratedFileStore()
    return pipeline


@pytest.mark.asyncio
async def test_extractor_handles_text_image_docx_and_invalid_extension():
    extractor = DocumentExtractor(FakeOCRAI())

    text_result = await extractor.extract(
        'note.txt', b'\xef\xbb\xbfHello world'
    )
    assert text_result.text == 'Hello world'

    latin_result = await extractor.extract('latin.txt', b'ol\xe1')
    assert latin_result.text == 'olá'

    image_result = await extractor.extract('scan.png', b'fake-image')
    assert image_result.text == 'OCR extracted text'
    assert len(image_result.visual_candidates) == 1
    assert image_result.visual_candidates[0].mime_type == 'image/png'

    document = Document()
    document.add_paragraph('Main paragraph')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'A'
    table.cell(0, 1).text = 'B'
    buffer = BytesIO()
    document.save(buffer)

    docx_result = await extractor.extract('report.docx', buffer.getvalue())
    assert 'Main paragraph' in docx_result.text
    assert 'A | B' in docx_result.text

    with pytest.raises(ValueError, match='Formato não suportado'):
        await extractor.extract('data.csv', b'value')


@pytest.mark.asyncio
async def test_extractor_handles_native_and_ocr_pdf_pages():
    extractor = DocumentExtractor(FakeOCRAI())

    native_pdf = fitz.open()
    page = native_pdf.new_page()
    page.insert_text(
        (72, 72),
        'Figura 1 chart with bars, legend, and enough native text for extraction.',
    )
    native_data = native_pdf.tobytes()
    native_pdf.close()

    native_result = await extractor.extract('native.pdf', native_data)
    assert 'Figura 1 chart' in native_result.text
    assert native_result.visual_candidates

    scanned_pdf = fitz.open()
    scanned_pdf.new_page()
    scanned_data = scanned_pdf.tobytes()
    scanned_pdf.close()

    scanned_result = await extractor.extract('scan.pdf', scanned_data)
    assert 'OCR extracted text' in scanned_result.text


def test_generated_file_store_creates_resolves_and_cleans_files(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(
        settings, 'ACCESSIBILITY_OUTPUT_DIR', str(tmp_path / 'output')
    )
    monkeypatch.setattr(settings, 'ACCESSIBILITY_FILE_TTL_SECONDS', 1)
    store = GeneratedFileStore()

    generated = store.new_path('.txt')
    generated.write_text('result', encoding='utf-8')
    assert generated.suffix == '.txt'
    assert store.resolve_safe(generated.name) == generated.resolve()

    with pytest.raises(ValueError, match='Nome de arquivo inválido'):
        store.resolve_safe('../outside.txt')

    old_file = store.root / 'old.txt'
    old_file.write_text('old', encoding='utf-8')
    old_timestamp = time.time() - 60
    os.utime(old_file, (old_timestamp, old_timestamp))
    store.cleanup()
    assert not old_file.exists()


@pytest.mark.asyncio
async def test_edge_tts_service_success_empty_and_invalid_output(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    class FakeCommunicate:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        async def save(self, path: str) -> None:
            Path(path).write_bytes(b'audio')

    monkeypatch.setattr(tts_module.edge_tts, 'Communicate', FakeCommunicate)
    service = TTSService()
    output = tmp_path / 'audio.mp3'
    await service.synthesize('Accessible audio text', output)
    assert output.read_bytes() == b'audio'

    with pytest.raises(ValueError, match='Não há texto'):
        await service.synthesize('   ', tmp_path / 'empty.mp3')

    class EmptyCommunicate:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        async def save(self, path: str) -> None:
            return None

    monkeypatch.setattr(tts_module.edge_tts, 'Communicate', EmptyCommunicate)
    with pytest.raises(RuntimeError):
        await service.synthesize('Text', tmp_path / 'missing.mp3')


def make_gemini_service(
    monkeypatch: pytest.MonkeyPatch,
    responses: list[object],
) -> GeminiService:
    class FakeModels:
        def __init__(self, items: list[object]) -> None:
            self.items = list(items)

        async def generate_content(self, **kwargs):
            assert kwargs['model'] == 'test-model'
            return self.items.pop(0)

    fake_client = SimpleNamespace(
        aio=SimpleNamespace(models=FakeModels(responses)),
    )
    monkeypatch.setattr(settings, 'GEMINI_API_KEY', 'test-key')
    monkeypatch.setattr(settings, 'GEMINI_MODEL', 'test-model')
    monkeypatch.setattr(gemini_module.genai, 'Client', lambda **_: fake_client)
    return GeminiService()


@pytest.mark.asyncio
async def test_gemini_service_requires_api_key(
    monkeypatch: pytest.MonkeyPatch,
):
    monkeypatch.setattr(settings, 'GEMINI_API_KEY', '')

    service = GeminiService()

    with pytest.raises(RuntimeError, match='GEMINI_API_KEY'):
        await service.generate_text(
            system_instruction='System',
            user_text='Input',
        )


@pytest.mark.asyncio
async def test_gemini_generate_text_and_ocr(monkeypatch: pytest.MonkeyPatch):
    service = make_gemini_service(
        monkeypatch,
        [
            SimpleNamespace(text='  generated text  '),
            SimpleNamespace(text='[SEM_TEXTO_LEGIVEL]'),
            SimpleNamespace(text='OCR text'),
        ],
    )

    result = await service.generate_text(
        system_instruction='System',
        user_text='Input',
    )
    assert result == 'generated text'

    assert not await service.ocr_image(b'image', 'image/png')
    assert await service.ocr_image(b'image', 'image/png') == 'OCR text'


@pytest.mark.asyncio
async def test_gemini_generate_text_rejects_empty_response(
    monkeypatch: pytest.MonkeyPatch,
):
    service = make_gemini_service(
        monkeypatch,
        [SimpleNamespace(text='   ')],
    )
    with pytest.raises(RuntimeError, match='resposta vazia'):
        await service.generate_text(
            system_instruction='System',
            user_text='Input',
        )


@pytest.mark.asyncio
async def test_gemini_chart_audit_and_correct(monkeypatch: pytest.MonkeyPatch):
    parsed_chart = ChartVisionResult(
        is_chart=True,
        title='Chart',
        description='Description',
        confidence=0.9,
    )
    parsed_audit = AuditReport(
        status='problemas_encontrados',
        itens=[AuditItem(problema='Missing value')],
    )
    service = make_gemini_service(
        monkeypatch,
        [
            SimpleNamespace(parsed=parsed_chart, text=''),
            SimpleNamespace(parsed=parsed_audit, text=''),
        ],
    )

    chart = await service.analyze_chart(
        image_bytes=b'image',
        mime_type='image/png',
        context_text='Nearby text',
    )
    assert chart.is_chart is True

    audit = await service.audit('Original', 'Accessible')
    assert audit.status == 'problemas_encontrados'
    assert audit.itens[0].problema == 'Missing value'

    service.generate_text = AsyncMock(return_value='Corrected text')
    corrected = await service.correct('Accessible', ['Issue one', 'Issue two'])
    assert corrected == 'Corrected text'
    call = service.generate_text.await_args.kwargs
    assert '1. Issue one' in call['user_text']
    assert '2. Issue two' in call['user_text']


@pytest.mark.asyncio
async def test_gemini_json_fallbacks(monkeypatch: pytest.MonkeyPatch):
    service = make_gemini_service(
        monkeypatch,
        [
            SimpleNamespace(
                parsed=None,
                text=(
                    '{"is_chart": true, "title": "Chart", '
                    '"description": "Trend", "confidence": 0.8}'
                ),
            ),
            SimpleNamespace(
                parsed=None,
                text='{"status": "ok", "itens": []}',
            ),
        ],
    )

    chart = await service.analyze_chart(
        image_bytes=b'image',
        mime_type='image/png',
        context_text='Context',
    )
    assert chart.description == 'Trend'

    audit = await service.audit('Original', 'Accessible')
    assert audit.status == 'ok'


@pytest.mark.asyncio
async def test_pipeline_success_with_math_chart_and_correction(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    extraction = ExtractionResult(
        text='x² + y² = 25. ∑ ∫ √ and additional mathematical context.',
        visual_candidates=[
            VisualCandidate(
                label='Page 1',
                image_bytes=b'chart',
                mime_type='image/png',
                context_text='Chart context',
            )
        ],
    )
    audit = AuditReport(
        status='problemas_encontrados',
        itens=[AuditItem(problema='Restore a missing value')],
    )
    ai = FakePipelineAI(audit=audit, corrected='Corrected accessible text')
    pipeline = make_pipeline(
        tmp_path=tmp_path,
        extraction=extraction,
        ai=ai,
        monkeypatch=monkeypatch,
    )

    events = [
        event
        async for event in pipeline.run(
            filename='math.pdf',
            data=b'pdf',
            level=2,
        )
    ]

    assert events[-1].type == 'result'
    result = events[-1].result
    assert result is not None
    assert result.accessible_text == 'Corrected accessible text'
    assert result.math_detection.is_math is True
    assert len(result.charts) == 1
    assert 'DESCRIÇÕES DE GRÁFICOS DETECTADOS' in ai.last_user_text
    assert 'ALÉM DAS REGRAS GERAIS' in ai.last_system_instruction
    assert (pipeline.store.root / Path(result.text_download_url).name).exists()
    assert (pipeline.store.root / Path(result.audio_url).name).exists()


@pytest.mark.asyncio
async def test_pipeline_success_without_math_chart_or_correction(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    extraction = ExtractionResult(text='Plain accessible document text.')
    ai = FakePipelineAI(audit=AuditReport(status='ok', itens=[]))
    pipeline = make_pipeline(
        tmp_path=tmp_path,
        extraction=extraction,
        ai=ai,
        monkeypatch=monkeypatch,
    )

    events = [
        event
        async for event in pipeline.run(
            filename='plain.txt',
            data=b'plain',
            level=1,
        )
    ]

    assert events[-1].type == 'result'
    result = events[-1].result
    assert result is not None
    assert result.charts == []
    assert result.audit.status == 'ok'


@pytest.mark.asyncio
async def test_pipeline_uses_visual_placeholder_when_no_text(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    extraction = ExtractionResult(
        text='',
        visual_candidates=[
            VisualCandidate(
                label='Uploaded image',
                image_bytes=b'chart',
                mime_type='image/png',
                context_text='visual',
            )
        ],
    )
    ai = FakePipelineAI(audit=AuditReport(status='ok', itens=[]))
    pipeline = make_pipeline(
        tmp_path=tmp_path,
        extraction=extraction,
        ai=ai,
        monkeypatch=monkeypatch,
    )

    events = [
        event
        async for event in pipeline.run(
            filename='chart.png',
            data=b'image',
            level=3,
        )
    ]

    assert events[-1].type == 'result'
    assert events[-1].result is not None
    assert events[-1].result.raw_text.startswith('[')


@pytest.mark.asyncio
async def test_pipeline_emits_error_for_empty_document(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    extraction = ExtractionResult(text='', visual_candidates=[])
    ai = FakePipelineAI(audit=AuditReport(status='ok', itens=[]))
    pipeline = make_pipeline(
        tmp_path=tmp_path,
        extraction=extraction,
        ai=ai,
        monkeypatch=monkeypatch,
    )

    events = [
        event
        async for event in pipeline.run(
            filename='empty.txt',
            data=b'empty',
            level=2,
        )
    ]

    assert events[-1].type == 'error'
    assert events[-1].message


@pytest.mark.asyncio
async def test_pipeline_rejects_invalid_level(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    extraction = ExtractionResult(text='Text')
    ai = FakePipelineAI(audit=AuditReport(status='ok', itens=[]))
    pipeline = make_pipeline(
        tmp_path=tmp_path,
        extraction=extraction,
        ai=ai,
        monkeypatch=monkeypatch,
    )

    events = [
        event
        async for event in pipeline.run(
            filename='document.txt',
            data=b'data',
            level=9,
        )
    ]

    assert len(events) == 1
    assert events[0].type == 'error'
    assert 'Nível inválido' in (events[0].message or '')


@pytest.mark.parametrize(
    ('style', 'expected'),
    [
        ('Heading 1', '# Derivadas'),
        ('Heading 3', '### Derivadas'),
        ('Título 2', '## Derivadas'),
        ('titulo 4', '#### Derivadas'),
        ('Title', '# Derivadas'),
        ('Normal', 'Derivadas'),
        (None, 'Derivadas'),
    ],
)
def test_docx_heading_styles_become_markdown(style, expected):
    paragraph = SimpleNamespace(style=SimpleNamespace(name=style))
    assert DocumentExtractor._markdown_heading(paragraph, 'Derivadas') == (
        expected
    )


def test_gemini_client_retries_transient_errors(
    monkeypatch: pytest.MonkeyPatch,
):
    created = {}
    monkeypatch.setattr(
        gemini_module.genai, 'Client', lambda **kwargs: created.update(kwargs)
    )

    GeminiService(api_key='test-key')

    retry = created['http_options'].retry_options
    assert retry.attempts > 1
    assert {429, 503} <= set(retry.http_status_codes)
