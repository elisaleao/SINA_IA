import json
import os
import time
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import docx
import fitz
import pytest
from httpx import AsyncClient

from app.api.routers.accessibility import (
    get_accessibility_pipeline,
    get_generated_file_store,
)
from app.main import app
from app.schemas.accessibility import (
    AuditItem,
    AuditReport,
    ChartVisionResult,
    PipelineEvent,
    ProcessResult,
)
from app.services.accessibility_pipeline import AccessibilityPipeline
from app.services.document_extractor import (
    DocumentExtractor,
    ExtractionResult,
    VisualCandidate,
)
from app.services.file_store import GeneratedFileStore
from app.services.gemini_service import GeminiService
from app.services.math_detector import detect_math_content
from app.services.tts_service import TTSService

# ---------------------------------------------------------------------------
# 1. Tests for GeneratedFileStore
# ---------------------------------------------------------------------------


def test_generated_file_store_lifecycle(tmp_path: Path):
    store = GeneratedFileStore(root_dir=tmp_path)
    assert store.root == tmp_path

    # new_path
    path = store.new_path('.txt')
    assert path.suffix == '.txt'
    assert path.parent == tmp_path
    path.write_text('teste', encoding='utf-8')

    # resolve_safe valid
    resolved = store.resolve_safe(path.name)
    assert resolved == path

    # resolve_safe invalid filenames
    with pytest.raises(ValueError, match='Nome de arquivo inválido.'):
        store.resolve_safe('')

    with pytest.raises(ValueError, match='Nome de arquivo inválido.'):
        store.resolve_safe('../etc/passwd')

    with pytest.raises(ValueError, match='Nome de arquivo inválido.'):
        store.resolve_safe('sub/dir/file.txt')


def test_generated_file_store_cleanup(tmp_path: Path):
    store = GeneratedFileStore(root_dir=tmp_path)
    store.max_age_seconds = 10

    old_file = tmp_path / 'old.txt'
    old_file.write_text('velho', encoding='utf-8')
    # Modifica o mtime para o passado
    old_time = time.time() - 20
    os.utime(old_file, (old_time, old_time))

    new_file = tmp_path / 'new.txt'
    new_file.write_text('novo', encoding='utf-8')

    store.cleanup()

    assert not old_file.exists()
    assert new_file.exists()


# ---------------------------------------------------------------------------
# 2. Tests for GeminiService
# ---------------------------------------------------------------------------


def test_gemini_service_initialization():
    service = GeminiService(api_key='fake-key', model='test-model')
    assert service.model == 'test-model'
    assert service.client is not None

    empty_service = GeminiService(api_key='', client=None)
    with patch('app.services.gemini_service.settings.GEMINI_API_KEY', ''):
        empty_service.client = None
        with pytest.raises(
            RuntimeError, match='GEMINI_API_KEY não configurada'
        ):
            empty_service._ensure_client()


@pytest.mark.asyncio
async def test_gemini_service_generate_text():
    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.text = 'Texto gerado acessível'
    mock_client.aio.models.generate_content = AsyncMock(return_value=mock_resp)

    service = GeminiService(client=mock_client)
    res = await service.generate_text(
        system_instruction='Instrução', user_text='Original'
    )
    assert res == 'Texto gerado acessível'

    # Empty response raises RuntimeError
    mock_resp.text = '   '
    with pytest.raises(RuntimeError, match='resposta vazia'):
        await service.generate_text(
            system_instruction='Instrução', user_text='Original'
        )


@pytest.mark.asyncio
async def test_gemini_service_ocr_image():
    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.text = 'Texto extraído da imagem'
    mock_client.aio.models.generate_content = AsyncMock(return_value=mock_resp)

    service = GeminiService(client=mock_client)
    res = await service.ocr_image(b'fake-image-bytes', 'image/png')
    assert res == 'Texto extraído da imagem'

    mock_resp.text = '[SEM_TEXTO_LEGIVEL]'
    res_empty = await service.ocr_image(b'fake-image-bytes', 'image/png')
    assert not res_empty


@pytest.mark.asyncio
async def test_gemini_service_analyze_chart():
    mock_client = MagicMock()
    mock_resp = MagicMock()
    chart_res = ChartVisionResult(
        is_chart=True,
        title='Gráfico de Vendas',
        description='Gráfico de barras mostrando vendas crescentes.',
        confidence=0.95,
    )
    mock_resp.parsed = chart_res
    mock_client.aio.models.generate_content = AsyncMock(return_value=mock_resp)

    service = GeminiService(client=mock_client)
    res = await service.analyze_chart(
        image_bytes=b'chart-bytes',
        mime_type='image/png',
        context_text='Figura 1: Vendas',
    )
    assert res.is_chart is True
    assert res.title == 'Gráfico de Vendas'

    # Fallback to json parsing when parsed is None
    mock_resp.parsed = None
    mock_resp.text = json.dumps({
        'is_chart': False,
        'title': None,
        'description': '',
        'confidence': 0.1,
    })
    res_fallback = await service.analyze_chart(
        image_bytes=b'chart-bytes',
        mime_type='image/png',
        context_text='Figura 1: Vendas',
    )
    assert res_fallback.is_chart is False


@pytest.mark.asyncio
async def test_gemini_service_audit_and_correct():
    mock_client = MagicMock()
    mock_resp = MagicMock()
    report = AuditReport(
        status='problemas_encontrados',
        itens=[AuditItem(problema='Faltou explicar a fórmula')],
    )
    mock_resp.parsed = report
    mock_client.aio.models.generate_content = AsyncMock(return_value=mock_resp)

    service = GeminiService(client=mock_client)
    audit = await service.audit('original', 'acessivel')
    assert audit.status == 'problemas_encontrados'
    assert len(audit.itens) == 1

    # Test correction
    mock_resp.text = 'Texto corrigido com a fórmula explicada'
    corrected = await service.correct(
        'acessivel', ['Faltou explicar a fórmula']
    )
    assert corrected == 'Texto corrigido com a fórmula explicada'


# ---------------------------------------------------------------------------
# 3. Tests for TTSService
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_edge_tts_empty_text(tmp_path: Path):
    tts = TTSService()
    with pytest.raises(ValueError, match='Não há texto para gerar áudio.'):
        await tts.synthesize('   ', tmp_path / 'out.mp3')


@pytest.mark.asyncio
async def test_edge_tts_synthesize_success(tmp_path: Path):
    tts = TTSService()
    out_file = tmp_path / 'out.mp3'

    async def fake_save(path_str):
        Path(path_str).write_bytes(b'ID3FakeAudioData')

    with patch('edge_tts.Communicate') as mock_communicate:
        instance = MagicMock()
        instance.save = AsyncMock(side_effect=fake_save)
        mock_communicate.return_value = instance

        await tts.synthesize('Texto para ler', out_file)
        assert out_file.exists()
        assert out_file.stat().st_size > 0


@pytest.mark.asyncio
async def test_edge_tts_synthesize_failure_raises(tmp_path: Path):
    tts = TTSService()
    out_file = tmp_path / 'out.mp3'

    with patch('edge_tts.Communicate') as mock_communicate:
        instance = MagicMock()
        instance.save = AsyncMock()  # doesn't write anything
        mock_communicate.return_value = instance

        with pytest.raises(
            RuntimeError, match='não gerou um arquivo de áudio válido'
        ):
            await tts.synthesize('Texto para ler', out_file)


# ---------------------------------------------------------------------------
# 4. Tests for DocumentExtractor
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_extractor_unsupported_format():
    ai = MagicMock(spec=GeminiService)
    extractor = DocumentExtractor(ai)
    with pytest.raises(ValueError, match='Formato não suportado'):
        await extractor.extract('planilha.xlsx', b'dummy-data')


@pytest.mark.asyncio
async def test_extractor_txt():
    ai = MagicMock(spec=GeminiService)
    extractor = DocumentExtractor(ai)
    content = 'Olá mundo! Este é um texto de teste em UTF-8.'
    res = await extractor.extract('teste.txt', content.encode('utf-8'))
    assert res.text == content
    assert len(res.visual_candidates) == 0


@pytest.mark.asyncio
async def test_extractor_docx():
    ai = MagicMock(spec=GeminiService)
    extractor = DocumentExtractor(ai)

    # Cria um DOCX na memória
    doc = docx.Document()
    doc.add_paragraph('Primeiro parágrafo do DOCX.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'A1'
    table.cell(0, 1).text = 'B1'
    table.cell(1, 0).text = 'A2'
    table.cell(1, 1).text = 'B2'

    bio = BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    res = await extractor.extract('documento.docx', docx_bytes)
    assert 'Primeiro parágrafo do DOCX.' in res.text
    assert '[Tabela 1]' in res.text
    assert 'A1 | B1' in res.text


@pytest.mark.asyncio
async def test_extractor_pdf_with_native_text():
    ai = MagicMock(spec=GeminiService)
    extractor = DocumentExtractor(ai)

    # Cria PDF com texto nativo usando PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (50, 50),
        'Texto nativo suficientemente longo para não disparar OCR de página inteira.',
    )
    pdf_bytes = doc.tobytes()
    doc.close()

    res = await extractor.extract('documento.pdf', pdf_bytes)
    assert '[Página 1]' in res.text
    assert 'Texto nativo suficientemente longo' in res.text


@pytest.mark.asyncio
async def test_extractor_pdf_with_short_text_triggers_ocr():
    ai = MagicMock(spec=GeminiService)
    ai.ocr_image = AsyncMock(return_value='Texto vindo de OCR')
    extractor = DocumentExtractor(ai)

    # Cria PDF quase vazio (<25 chars)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), 'Curto')
    pdf_bytes = doc.tobytes()
    doc.close()

    res = await extractor.extract('scaneado.pdf', pdf_bytes)
    assert 'Texto vindo de OCR' in res.text
    ai.ocr_image.assert_awaited_once()


@pytest.mark.asyncio
async def test_extractor_image():
    ai = MagicMock(spec=GeminiService)
    ai.ocr_image = AsyncMock(return_value='Texto reconhecido da imagem PNG')
    extractor = DocumentExtractor(ai)

    res = await extractor.extract('sample_image.png', b'fake-png-data')
    assert res.text == 'Texto reconhecido da imagem PNG'
    assert len(res.visual_candidates) == 1
    assert res.visual_candidates[0].mime_type == 'image/png'


# ---------------------------------------------------------------------------
# 5. Tests for AccessibilityPipeline
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_accessibility_pipeline_invalid_level():
    pipeline = AccessibilityPipeline()
    events = []
    async for event in pipeline.run(
        filename='teste.txt', data=b'ola', level=99
    ):
        events.append(event)
    assert len(events) == 1
    assert events[0].type == 'error'
    assert 'Nível inválido' in (events[0].message or '')


@pytest.mark.asyncio
async def test_accessibility_pipeline_empty_content_error():
    ai = MagicMock(spec=GeminiService)
    extractor = MagicMock(spec=DocumentExtractor)
    extractor.extract = AsyncMock(
        return_value=ExtractionResult(text='', visual_candidates=[])
    )

    pipeline = AccessibilityPipeline(ai=ai, extractor=extractor)
    events = []
    async for event in pipeline.run(filename='teste.txt', data=b' ', level=2):
        events.append(event)

    error_events = [e for e in events if e.type == 'error']
    assert len(error_events) == 1
    assert 'Não foi possível extrair conteúdo' in (
        error_events[0].message or ''
    )


@pytest.mark.asyncio
async def test_accessibility_pipeline_full_run_with_audit_correction(
    tmp_path: Path,
):
    ai = MagicMock(spec=GeminiService)
    ai.generate_text = AsyncMock(return_value='Texto acessível gerado.')
    ai.analyze_chart = AsyncMock(
        return_value=ChartVisionResult(
            is_chart=True,
            title='Gráfico X',
            description='Descrição do gráfico',
            confidence=0.9,
        )
    )
    ai.audit = AsyncMock(
        return_value=AuditReport(
            status='problemas_encontrados',
            itens=[AuditItem(problema='Ajustar leitura de fração')],
        )
    )
    ai.correct = AsyncMock(
        return_value='Texto acessível devidamente corrigido.'
    )

    tts = MagicMock(spec=TTSService)

    async def fake_synth(text, path):
        Path(path).write_bytes(b'fake-mp3')

    tts.synthesize = AsyncMock(side_effect=fake_synth)

    store = GeneratedFileStore(root_dir=tmp_path)

    extractor = MagicMock(spec=DocumentExtractor)
    extractor.extract = AsyncMock(
        return_value=ExtractionResult(
            text='Documento com matemática: x² = 4 e gráfico.',
            visual_candidates=[
                VisualCandidate(
                    label='Figura 1',
                    image_bytes=b'img',
                    mime_type='image/png',
                    context_text='contexto',
                )
            ],
        )
    )

    pipeline = AccessibilityPipeline(
        ai=ai, tts=tts, store=store, extractor=extractor
    )

    events: list[PipelineEvent] = []
    async for event in pipeline.run(
        filename='artigo.pdf', data=b'pdfbytes', level=2
    ):
        events.append(event)

    result_event = next((e for e in events if e.type == 'result'), None)
    assert result_event is not None
    assert result_event.result is not None
    assert result_event.result.filename == 'artigo.pdf'
    assert (
        result_event.result.accessible_text
        == 'Texto acessível devidamente corrigido.'
    )
    assert len(result_event.result.charts) == 1
    assert result_event.result.charts[0].title == 'Gráfico X'
    assert result_event.result.audit.status == 'problemas_encontrados'


# ---------------------------------------------------------------------------
# 6. Tests for Router Endpoints (/api/accessibility)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_router_process_stream_validation_errors(client: AsyncClient):
    # Formato não suportado
    resp = await client.post(
        '/api/accessibility/process-stream',
        files={
            'file': ('teste.exe', b'fake binary', 'application/octet-stream')
        },
        data={'level': 2},
    )
    assert resp.status_code == 415

    # Nível inválido
    resp = await client.post(
        '/api/accessibility/process-stream',
        files={'file': ('teste.txt', b'texto', 'text/plain')},
        data={'level': 10},
    )
    assert resp.status_code == 422

    # Arquivo vazio
    resp = await client.post(
        '/api/accessibility/process-stream',
        files={'file': ('teste.txt', b'', 'text/plain')},
        data={'level': 2},
    )
    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_router_process_stream_success(client: AsyncClient):
    mock_pipeline = MagicMock(spec=AccessibilityPipeline)

    async def mock_run(*args, **kwargs):
        yield PipelineEvent(
            type='stage', stage='upload', status='done', message='Ok'
        )
        yield PipelineEvent(
            type='result',
            result=ProcessResult(
                filename='teste.txt',
                level=2,
                raw_text='texto',
                accessible_text='acessivel',
                math_detection=detect_math_content('texto'),
                charts=[],
                audit=AuditReport(status='ok', itens=[]),
                text_download_url='/api/accessibility/files/abc.txt',
                audio_url='/api/accessibility/files/abc.mp3',
            ),
        )

    mock_pipeline.run = mock_run
    app.dependency_overrides[get_accessibility_pipeline] = lambda: (
        mock_pipeline
    )

    try:
        resp = await client.post(
            '/api/accessibility/process-stream',
            files={'file': ('teste.txt', b'Ola mundo', 'text/plain')},
            data={'level': 2},
        )
        assert resp.status_code == 200
        assert 'application/x-ndjson' in resp.headers.get('content-type', '')
        lines = [line for line in resp.text.strip().split('\n') if line]
        assert len(lines) == 2
        ev1 = json.loads(lines[0])
        assert ev1['type'] == 'stage'
        assert ev1['stage'] == 'upload'
        ev2 = json.loads(lines[1])
        assert ev2['type'] == 'result'
        assert ev2['result']['accessible_text'] == 'acessivel'
    finally:
        app.dependency_overrides.pop(get_accessibility_pipeline, None)


@pytest.mark.asyncio
async def test_router_download_files(client: AsyncClient, tmp_path: Path):
    store = GeneratedFileStore(root_dir=tmp_path)
    txt_file = tmp_path / 'resultado.txt'
    txt_file.write_text('Texto acessivel gerado', encoding='utf-8')
    mp3_file = tmp_path / 'audio.mp3'
    mp3_file.write_bytes(b'FakeMp3Audio')

    app.dependency_overrides[get_generated_file_store] = lambda: store

    try:
        # Download TXT existente
        resp = await client.get('/api/accessibility/files/resultado.txt')
        assert resp.status_code == 200
        assert 'text/plain' in resp.headers.get('content-type', '')
        assert 'Texto acessivel gerado' in resp.text

        # Download MP3 existente
        resp_audio = await client.get('/api/accessibility/files/audio.mp3')
        assert resp_audio.status_code == 200
        assert 'audio/mpeg' in resp_audio.headers.get('content-type', '')

        # Arquivo inexistente
        resp_404 = await client.get('/api/accessibility/files/inexistente.txt')
        assert resp_404.status_code == 404

        # Path traversal / invalid filename
        resp_400 = await client.get('/api/accessibility/files/%2E%2E')
        assert resp_400.status_code == 400
    finally:
        app.dependency_overrides.pop(get_generated_file_store, None)
