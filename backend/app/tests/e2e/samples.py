import io

import docx
import pymupdf

SAMPLE_TEX = r"""\documentclass{article}
\usepackage{amsmath}
\begin{document}
\section{Relatividade}
A energia de repouso é $E = mc^2$ e Pitágoras diz $a^2 + b^2 = c^2$.
\begin{equation}
F = m a
\end{equation}
\end{document}
""".encode()

CORRUPTED_PDF = b'%PDF-1.7\n1 0 obj << /Type /Catalog ' + b'\xff' * 64

MALICIOUS_EXE = b'MZ\x90\x00\x03\x00\x00\x00' + b'\x00' * 120


def sample_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    lines = [
        'Cinemática',
        'Velocidade média é deslocamento dividido pelo tempo.',
        'Grandeza | Unidade',
        'Velocidade | m/s',
        'Aceleração | m/s²',
    ]
    for index, line in enumerate(lines):
        page.insert_text((72, 72 + index * 20), line)
    data = document.tobytes()
    document.close()
    return data


def sample_docx() -> bytes:
    document = docx.Document()
    document.add_heading('Leis de Newton', level=1)
    document.add_paragraph(
        'Ação e reação: forças têm o mesmo módulo e sentidos opostos.'
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
