from __future__ import annotations

import asyncio
import mimetypes
import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

from .gemini_service import GeminiAccessibilityService

SUPPORTED_EXTENSIONS = {
    '.pdf',
    '.docx',
    '.txt',
    '.png',
    '.jpg',
    '.jpeg',
    '.webp',
}
_CHART_WORDS = re.compile(
    r'\b(gr[aá]fico|figura|eixo|eixos|histograma|dispers[aã]o|boxplot|'
    r'barras?|linhas?|pizza|percentual|porcentagem|legenda|s[eé]rie)\b',
    flags=re.IGNORECASE,
)


@dataclass(slots=True)
class VisualCandidate:
    label: str
    image_bytes: bytes
    mime_type: str
    context_text: str


@dataclass(slots=True)
class _PdfPageScan:
    number: int
    native_text: str
    has_visual_hint: bool
    ocr_render: bytes | None
    chart_render: bytes | None


@dataclass(slots=True)
class ExtractionResult:
    text: str
    visual_candidates: list[VisualCandidate] = field(default_factory=list)


class DocumentExtractor:
    """
    Extrai texto no backend e só usa Gemini como OCR quando o documento não
    oferece texto nativo suficiente. Isso reduz chamadas externas.
    """

    def __init__(
        self,
        ai: GeminiAccessibilityService,
        *,
        max_visual_candidates: int = 4,
    ) -> None:
        self.ai = ai
        self.max_visual_candidates = max_visual_candidates

    async def extract(self, filename: str, data: bytes) -> ExtractionResult:
        extension = Path(filename).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                'Formato não suportado. Use PDF, DOCX, TXT, PNG, JPG/JPEG ou WEBP.'
            )

        if extension == '.txt':
            return ExtractionResult(text=self._decode_text(data))
        if extension == '.docx':
            # python-docx e zipfile são CPU: fora do event loop
            return await asyncio.to_thread(self._extract_docx, data)
        if extension == '.pdf':
            return await self._extract_pdf(data)
        return await self._extract_image(extension, data)

    @staticmethod
    def _decode_text(data: bytes) -> str:
        for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
            try:
                return data.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        return data.decode('utf-8', errors='replace').strip()

    def _extract_docx(self, data: bytes) -> ExtractionResult:
        document = Document(BytesIO(data))
        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table_index, table in enumerate(document.tables, start=1):
            parts.append(f'[Tabela {table_index}]')
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                if row_text.strip(' |'):
                    parts.append(row_text)

        text = '\n'.join(parts).strip()
        candidates: list[VisualCandidate] = []
        context = text[:6000]

        with zipfile.ZipFile(BytesIO(data)) as archive:
            media_names = sorted(
                name
                for name in archive.namelist()
                if name.startswith('word/media/') and not name.endswith('/')
            )
            for index, media_name in enumerate(
                media_names[: self.max_visual_candidates], start=1
            ):
                image_bytes = archive.read(media_name)
                extension = Path(media_name).suffix.lower()
                mime_type = mimetypes.types_map.get(extension, 'image/png')
                if not mime_type.startswith('image/'):
                    continue
                candidates.append(
                    VisualCandidate(
                        label=f'Imagem {index} do DOCX',
                        image_bytes=image_bytes,
                        mime_type=mime_type,
                        context_text=context,
                    )
                )

        return ExtractionResult(text=text, visual_candidates=candidates)

    async def _extract_pdf(self, data: bytes) -> ExtractionResult:
        # PyMuPDF roda numa thread; só as chamadas de OCR ficam no event loop
        pages = await asyncio.to_thread(self._scan_pdf, data)
        page_texts: list[str] = []
        candidates: list[VisualCandidate] = []

        for page in pages:
            native_text = page.native_text
            # OCR apenas quando há pouco texto selecionável.
            if page.ocr_render is not None:
                native_text = (
                    await self.ai.ocr_image(page.ocr_render, 'image/png')
                ).strip()

            page_texts.append(f'[Página {page.number}]\n{native_text}'.strip())

            if len(candidates) >= self.max_visual_candidates:
                continue

            likely_chart = page.has_visual_hint or bool(
                _CHART_WORDS.search(native_text)
            )
            rendered = page.ocr_render or page.chart_render
            if likely_chart and rendered is not None:
                candidates.append(
                    VisualCandidate(
                        label=f'Página {page.number}',
                        image_bytes=rendered,
                        mime_type='image/png',
                        context_text=native_text[:6000],
                    )
                )

        return ExtractionResult(
            text='\n\n'.join(page_texts).strip(),
            visual_candidates=candidates,
        )

    def _scan_pdf(self, data: bytes) -> list[_PdfPageScan]:
        pdf = fitz.open(stream=data, filetype='pdf')
        pages: list[_PdfPageScan] = []
        chart_renders = 0
        try:
            for page_number, page in enumerate(pdf, start=1):
                native_text = page.get_text('text').strip()
                image_count = len(page.get_images(full=True))
                drawing_count = len(page.get_drawings())
                has_visual_hint = drawing_count >= 12 or image_count > 0

                ocr_render: bytes | None = None
                chart_render: bytes | None = None
                if len(native_text) < 25:
                    ocr_render = self._render_page(page, scale=2.0)
                elif chart_renders < self.max_visual_candidates and (
                    has_visual_hint or _CHART_WORDS.search(native_text)
                ):
                    chart_render = self._render_page(page, scale=1.6)
                    chart_renders += 1

                pages.append(
                    _PdfPageScan(
                        number=page_number,
                        native_text=native_text,
                        has_visual_hint=has_visual_hint,
                        ocr_render=ocr_render,
                        chart_render=chart_render,
                    )
                )
        finally:
            pdf.close()
        return pages

    async def _extract_image(
        self, extension: str, data: bytes
    ) -> ExtractionResult:
        mime_type = {
            '.png': 'image/png',
            '.jpg': 'image/jpeg',
            '.jpeg': 'image/jpeg',
            '.webp': 'image/webp',
        }[extension]
        text = (await self.ai.ocr_image(data, mime_type)).strip()
        return ExtractionResult(
            text=text,
            visual_candidates=[
                VisualCandidate(
                    label='Imagem enviada',
                    image_bytes=data,
                    mime_type=mime_type,
                    context_text=text[:6000],
                )
            ],
        )

    @staticmethod
    def _render_page(page: fitz.Page, *, scale: float) -> bytes:
        matrix = fitz.Matrix(scale, scale)
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        return pixmap.tobytes('png')
